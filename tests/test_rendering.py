"""Regression tests for the document rendering contract.

These cover the defects that made generated resumes wrong:
  1. The template's variable names did not match the render context, so the
     candidate name, title, profile and institution rendered blank.
  2. docxtpl renders with autoescape=False by default, which dropped literal
     ampersands ('Data & AI' became 'Data  AI').
  3. The experience section repeated the company as a "Client / Project" line,
     because the schema nested projects the model filled with company names.
"""
import io
import os

import pytest

os.environ.setdefault("ALLOW_ANONYMOUS_AUTH", "true")

from docx import Document  # noqa: E402

from src.app import (  # noqa: E402
    TEMPLATE_VARIABLES,
    EducationItem,
    _compose_institution_line,
    _extract_json_object,
    build_render_context,
    render_resume,
    sanitize_resume_dict,
    template_contract_report,
)

SAMPLE = {
    "full_name": "RAVI INALA",
    "headline": "Architect - Data & AI",
    "profile_summary": "Architect across Cloud & AI and BI & Reporting.",
    "skills": ["Azure Databricks", "ETL & Data Modelling"],
    "education": [{"degree": "M.Tech Computer Science", "institution": "IIT Delhi", "score": "8.5 CGPA"}],
    "certifications": ["Azure Solutions Architect Expert"],
    "experience_summary_bullets": ["Led 20+ Data & AI platform builds."],
    "work_experience": [
        {
            "company_name": "Insight Direct",
            "overall_title": "Lead Architect",
            "duration": "Apr 2012 - Present",
            "achievements": ["Cut ETL runtime 60% for Microsoft."],
        }
    ],
}


def _rendered_text(payload: dict) -> str:
    doc = Document(io.BytesIO(render_resume(build_render_context(payload))))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_template_contract_is_satisfied():
    """Every tag in the .docx must be supplied by build_render_context."""
    report = template_contract_report()
    assert report["ok"], f"Template expects variables we never supply: {report.get('unsatisfied')}"
    assert set(report["template_variables"]) <= TEMPLATE_VARIABLES


def test_context_uses_template_variable_names():
    ctx = build_render_context(SAMPLE)
    assert ctx["name"] == "RAVI INALA"
    assert ctx["title"] == "Architect - Data & AI"
    assert ctx["profile"].startswith("Architect across")
    assert ctx["education"][0]["institution_line"] == "IIT Delhi | 8.5 CGPA"


@pytest.mark.parametrize("field", ["RAVI INALA", "PROFILE", "M.Tech Computer Science", "IIT Delhi | 8.5 CGPA"])
def test_core_fields_actually_render(field):
    """Regression: these rendered blank because of the name mismatch."""
    assert field in _rendered_text(SAMPLE)


def test_ampersands_survive_rendering():
    """Regression: autoescape=False silently ate every literal '&'."""
    text = _rendered_text(SAMPLE)
    assert "Data & AI" in text
    assert "ETL & Data Modelling" in text
    assert "Data  AI" not in text


def test_xml_metacharacters_do_not_corrupt_the_document():
    payload = dict(SAMPLE, full_name='Jane <script> & "Doe"')
    assert 'Jane <script> & "Doe"' in _rendered_text(payload)


# --------------------------------------------- company-centric experience
def test_company_role_and_dates_render_on_their_own_lines():
    text = _rendered_text(SAMPLE)
    assert "Insight Direct\tApr 2012 - Present" in text   # one line, tabbed to the margin
    assert "Lead Architect" in text
    assert "Cut ETL runtime 60% for Microsoft." in text


def test_no_duplicated_client_project_line():
    """Regression: the old schema emitted 'Client / Project: <company> - <role>'."""
    text = _rendered_text(SAMPLE)
    assert "Client / Project" not in text


def test_company_without_a_role_renders_cleanly():
    payload = dict(SAMPLE, work_experience=[
        {"company_name": "Acme", "overall_title": "", "duration": "2019 - 2020",
         "achievements": ["Did a thing."]}
    ])
    text = _rendered_text(payload)
    assert "Acme\t2019 - 2020" in text
    assert "Did a thing." in text


def test_missing_optional_data_renders_without_error():
    """StrictUndefined must not trip on legitimately empty sections."""
    sparse = {"full_name": "A B", "headline": "", "profile_summary": ""}
    assert "A B" in _rendered_text(sparse)


def test_compose_institution_line_without_score():
    assert _compose_institution_line("IIT Delhi", None) == "IIT Delhi"


def test_extract_json_object_strips_fences_and_prose():
    assert _extract_json_object('```json\n{"a": 1}\n```') == '{"a": 1}'
    assert _extract_json_object('Sure!\n{"a": {"b": 2}}\nHope that helps') == '{"a": {"b": 2}}'


def test_extract_json_object_ignores_braces_inside_strings():
    assert _extract_json_object('{"a": "a } brace"}') == '{"a": "a } brace"}'


# ------------------------------------------------ education without a school
def _education_block(payload: dict) -> list:
    """Paragraphs from the degree up to the next section heading."""
    doc = Document(io.BytesIO(render_resume(build_render_context(payload))))
    paras = [p.text.strip() for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    start = paras.index("B.Tech (Computer Science)")
    return paras[start:paras.index("CERTIFICATIONS", start)]


def test_degree_without_institution_renders_alone():
    """Regression: the institution line rendered 'Institution Name Not Provided'."""
    payload = sanitize_resume_dict(dict(SAMPLE, education=[
        {"degree": "B.Tech (Computer Science)", "institution": "Institution Name Not Provided", "score": None}
    ]))
    # Degree, then only the template's own spacer before CERTIFICATIONS: no
    # placeholder text and no blank line where the institution would have been.
    assert _education_block(payload) == ["B.Tech (Computer Science)", ""]


def test_institution_still_renders_when_present():
    payload = dict(SAMPLE, education=[
        {"degree": "B.Tech (Computer Science)", "institution": "JNTU Hyderabad", "score": ""}
    ])
    assert _education_block(payload) == ["B.Tech (Computer Science)", "JNTU Hyderabad", ""]


@pytest.mark.parametrize("placeholder", [
    "Institution Name Not Provided", "Not Provided", "N/A", "n/a", "Unknown",
    "Institution: Not specified", "(not mentioned)", "-", "None",
])
def test_placeholders_are_blanked(placeholder):
    cleaned = sanitize_resume_dict({"education": [{"degree": "MBA", "institution": placeholder}]})
    assert cleaned["education"][0]["institution"] == ""


@pytest.mark.parametrize("real", [
    "Institute of Unknown Studies", "Not Provided by HR: escalated access requests",
    "National Academy", "Nanyang Technological University",
])
def test_real_values_are_not_mistaken_for_placeholders(real):
    cleaned = sanitize_resume_dict({"education": [{"degree": "MBA", "institution": real}]})
    assert cleaned["education"][0]["institution"] == real


def test_placeholder_list_items_are_dropped():
    cleaned = sanitize_resume_dict({"certifications": ["AZ-305", "N/A"]})
    assert cleaned["certifications"] == ["AZ-305"]


def test_schema_does_not_force_an_institution():
    assert EducationItem(degree="B.Tech").institution == ""
